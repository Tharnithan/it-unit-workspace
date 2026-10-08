from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

d=Document()
s=d.sections[0]; s.page_height=Inches(11.7); s.page_width=Inches(8.3)
s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(26)
d.styles['Heading 1'].font.size=Pt(18)
d.styles['Heading 2'].font.size=Pt(13)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def bullets(items):
 for t in items: d.add_paragraph(t,style='List Bullet')
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for c,x in zip(t.rows[0].cells,headers): c.text=x
 rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
 for row in rows:
  for c,x in zip(t.add_row().cells,row): c.text=x
 for i,row in enumerate(t.rows):
  for c,w in zip(row.cells,widths):
   c.width=Inches(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   pr.append(borders)
   sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7EDF3' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF')); pr.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.space_before=Pt(4); pp.paragraph_format.space_after=Pt(4)
    for r in pp.runs: r.font.size=Pt(10); r.bold=(i==0)
 p('')

d.add_paragraph('IT Division Work Management System',style='Title')
d.add_paragraph('Project Proposal and Implementation Plan',style='Subtitle')
p('Prepared for the Head of the IT Division and the IT team\nSri Lanka Institute of Technology\n8 October 2026 | Version 1.0')
h('Proposal summary')
p('Develop a secure internal website that brings IT work assignments, progress tracking, staff contacts, chat, meetings and notifications into one system. Mr. Kanchana, the SDD and head of the IT Division, will assign work and monitor delivery. Employees will update their assigned tasks, communicate with colleagues and view upcoming meetings. Completed work will remain in the database for monthly reports.')
p('The recommended solution is a responsive React and TypeScript website backed by Supabase PostgreSQL, authentication, storage and Realtime, with Firebase Cloud Messaging for web push. This architecture supports a free initial rollout within provider quotas and keeps task and reporting data in a relational database.')
h('Business objectives')
bullets(['Give the SDD a clear view of workload, overdue work and progress across the IT team.','Give every employee a personal work list and reliable record of assignments and updates.','Centralize team contact details, private conversations and the IT unit group chat.','Preserve completed tasks and produce consistent monthly reports by employee and supported division.'])
h('Scope of the first release')
p('The first release serves the IT Division only. Other divisions are recorded as the recipients of IT support; they do not receive user accounts in this phase. A future service request portal can let those divisions submit and track support requests.')
p('This proposal covers requirements, permissions, technical design, data structure, delivery stages and acceptance criteria. Implementation starts after the organization confirms the staff directory and operating policies.')

page('1 Organization and user roles')
p('The initial directory contains 11 named IT staff members, plus a separate administrator account. Job titles are directory information; application permissions are assigned separately as SDD, employee or administrator.')
table(['Member','Designation','Application role'],[
('Mr. Kanchana','SDD and head of IT','SDD'),('Mr. Insaf','TA','Employee'),('Mr. Damith','TA','Employee'),('Mr. Tharindu','Technician','Employee'),('Ms. Kinkini','Software Assistant','Employee'),('Ms. Ayali','MA','Employee'),('Mrs. Kumari','MA','Employee'),('Mr. Lakshitha','SE Intern','Employee'),('T. Tharnithan','SE Intern','Employee'),('Mr. Nisal','SE Intern','Employee'),('Mr. Sumeshan','Designation to be confirmed','Employee'),('System administrator','System maintenance','Administrator')],[2,2.8,1.9])
h('Supported divisions')
p('Scientific Standardization Division; Engineering Standardization Division; Product Certification Division; Systems Certification Division; Quality Assurance Division; Laboratory Services Division; Metrology Division; Standards & Services Promotion Division; Documentation & Information (D&I) Division.')
h('Staff directory fields')
p('Store full name, designation, application role, profile photo, official email, IT mailbox identifier such as itunit1, LAN phone extension, account status and joining or leaving dates. Do not invent an email domain or assign mailbox numbers without the existing staff mapping. Each login must belong to a named person even when a mailbox is shared.')

page('2 Features and access permissions')
table(['Action','Employee','SDD','Administrator'],[
('View staff contacts','Yes','Yes','Yes'),('Add edit or deactivate staff','No','Yes','Yes'),('Assign or reassign tasks','No','Yes','Yes'),('Update task progress','Own assignment','Any task','Any task'),('Complete a task','Own assignment','Any task','Override'),('View all task details','No by default','Yes','Yes'),('Schedule meetings','No by default','Yes','Yes'),('Private and unit chat','Participant','Participant','Moderation access'),('Monthly report export','Own work','All work','All work'),('Manage settings and deletion','No','Limited','Yes')],[2.9,1.2,1.2,1.4])
p('Recommended default: employees see their own tasks; the SDD and administrator see all tasks. Broader team task visibility can be enabled after the division confirms that policy. All staff can view ordinary work contact details, but only the SDD and administrator can edit directory records.')
h('Application sections')
bullets(['Dashboard and Works to Do: personal assignments, due dates, overdue items, progress and upcoming meetings; SDD dashboard includes team workload.','Tasks: title, description, supported division, assignee, priority, due date, status, percentage, progress notes, attachments and history.','Chat: direct conversations between two members and a persistent IT unit group; unread counts and message history.','Meetings: agenda, organizer, date, time, location or external meeting link, invitees, attendance response and minutes.','Team and Profiles: searchable staff directory with email and LAN extension.','Notifications: persistent inbox, unread badges, push preferences and links to the relevant record.','Reports and Administration: monthly exports, account management, audit history and maintenance settings.'])
p('Meeting scheduling is included. Built-in video calls, call recording, SMS, LAN telephone integration and a separate mobile app are outside the first release. A meeting can contain a link to an existing video meeting service.')

page('3 Task workflow and reporting')
h('Assignment and progress')
p('The SDD creates a task and chooses one responsible employee, the supported division, priority and due date. The employee receives a notification and starts the task. Progress updates contain status, percentage, a work note and optional evidence. Use Assigned, In Progress, Blocked and Completed as the core statuses; overdue is calculated from the due date, rather than stored as a status.')
h('Completion rule')
p('Either the SDD or the currently assigned employee may complete the task. No other employee may close it. Completion requires a resolution note and records completed_by and completed_at in one database transaction. Completing a task sets progress to 100 percent. Setting progress to 100 percent alone does not close it. This proposal does not require both people to approve completion.')
p('The administrator has an explicit maintenance override for correction or removal, as requested. Overrides require a reason and an audit event. Completed tasks become read-only for employees. The SDD or administrator can reopen a task with a reason; the original completion remains in the event history. Reassignment removes the previous assignee’s update and completion permission.')
h('Monthly report')
p('The SDD selects a month and exports a printable report and CSV. Report completion dates use Asia/Colombo time, with an inclusive month start and exclusive next-month start. Store timestamps in UTC. Report rows include task ID, title, supported division, employee at completion, created date, due date, completion date, priority and resolution.')
p('Summary measures include tasks created, completed, completed late, open at month end and blocked at month end. Reconstruct historical status from task events. Capture employee and division labels at completion so later directory edits do not rewrite past reports. Count distinct task IDs for monthly completions and list reopenings separately to avoid inflated totals.')
h('Retention and deletion')
p('Retain completed tasks in the same database with their updates and audit history; archiving is a visibility flag, not deletion. Deactivating an employee preserves past assignments and reporting records. The administrator can correct or remove records; normal removal uses soft deletion, while permanent purge is a separate maintenance action with a reason, reference checks and a backup. Agree retention periods before enabling purge.')
h('Example acceptance scenario')
p('Mr. Kanchana assigns a laboratory support task to Mr. Insaf. Mr. Insaf updates progress and closes it with a resolution. Mr. Damith cannot update or complete that task. The SDD sees the completion immediately and the task appears in the correct monthly report.')

page('4 Language and backend recommendation')
table(['Layer','Recommended technology','Purpose'],[
('Website','React and TypeScript with Vite','Responsive browser application'),('Interface styling','Tailwind CSS','Consistent layouts and components'),('Database','Supabase PostgreSQL and SQL','Structured records and reports'),('Login and authorization','Supabase Auth and row level security','Named accounts and data permissions'),('Live updates','Supabase Realtime','Chat and notification updates'),('Server logic','Supabase Edge Functions in TypeScript','Account administration and push sender'),('Files','Private Supabase Storage buckets','Task attachments and profile photos'),('Push notifications','Firebase Cloud Messaging','Background web notifications'),('Website hosting','Cloudflare Pages','Static site served over HTTPS')],[1.65,2.65,2.4])
p('TypeScript is the primary website language; HTML and CSS remain the foundations of the interface. SQL defines tables, permissions and reporting queries. React with Vite is appropriate for an internal application that does not require search engine indexing or a separate server-rendered website.')
h('Backend responsibilities')
p('Use database functions for controlled task transitions and transactions, rather than letting clients write arbitrary status changes. Row level security protects every exposed table. Use authenticated Edge Functions for privileged account operations and notification delivery; never expose Supabase service credentials or Firebase sender credentials in browser code.')
p('Normal flow: browser → Supabase Auth → authorized database queries or task functions → PostgreSQL. Realtime supplies live changes. Private storage policies authorize attachment downloads. Invite-only registration prevents public signups; the SDD or administrator creates employee records through server-side operations. Only the administrator can grant administrator privileges.')
h('Alternatives')
p('Firebase Authentication and Firestore are another option for live chat and push, but relational monthly reporting and task history need more deliberate denormalization. Laravel with PostgreSQL is suitable if the team prefers PHP; Django with PostgreSQL is suitable if it prefers Python. Both need a separately operated server, background jobs and a real-time chat solution. Self-hosted PostgreSQL or Supabase can use an existing institute server, but electricity, backups and maintenance remain operational costs.')

page('5 Database security and notifications')
h('Proposed database model')
table(['Entity','Main fields and relationships'],[
('profiles and divisions','Auth user ID, role, designation, contacts, active status; supported division directory'),('tasks','Assignee, division, creator, status, progress, priority, due date, completion actor and timestamp'),('task_updates and task_events','Task ID, actor, note, evidence; immutable assignment and status events with timestamps'),('conversations and participants','Direct or group type; membership determines access'),('messages','Conversation ID, sender, body, sent time and edit or deletion metadata'),('meetings and meeting_attendees','Organizer, agenda, times, location, link; invitee response and attendance'),('notifications and push_devices','Recipient, event, record link, read time; per-user per-device tokens'),('notification_outbox and audit_log','Delivery attempts, retry state, event key; privileged change actor and reason'),('attachments and report_runs','Private file paths and parent record; monthly export metadata and snapshot')],[2.25,4.45])
h('Authorization and operational security')
p('Enforce roles and ownership on the server, including database queries, real-time subscriptions and storage. Employees cannot change their own role or directory details. Chat messages are available only to conversation participants; the SDD does not automatically gain access to other people’s private chats. Administrator moderation access is exceptional and logged, and is explained to users. The proposed chat is not end-to-end encrypted.')
p('Require strong passwords and administrator MFA, expire sessions, validate input, rate-limit chat and uploads, restrict file types and sizes, and keep secrets in managed server configuration. Export the database and attachment inventory regularly to an institute-controlled location and test restoration. Maintain separate development and production environments.')
h('Reliable notification flow')
p('A task assignment, progress event, completion, chat message or meeting change writes a persistent notification and an outbox row in the same transaction. A server worker sends FCM messages to the recipient’s registered devices. A scheduled job retries failures with backoff, removes invalid tokens and uses an event key to prevent duplicate alerts. Scheduled reminders run in server jobs, including when no browser is open.')
p('The website requests notification permission after a user action, registers a service worker and stores device tokens under that user. Push requires HTTPS and browser support. Validate institute PCs and mobile browsers during the pilot. Permission refusal, offline devices or disabled browser background activity can prevent delivery; the persistent inbox remains available. Send a generic alert and record link rather than confidential chat text on lock screens. Read actions and links must recheck access.')

page('6 Free service plan and operating limits')
p('Target recurring service cost for a small pilot: USD 0 per month, within free allowances. Development time, institute devices, internet connectivity and any custom domain are separate costs. Free quotas can change; check the cited provider pages again before deployment.')
table(['Service','Free allowance or condition','Planning response'],[
('Supabase database and Auth','500 MB database; 50,000 monthly active users','More than sufficient account capacity for the initial team; monitor database growth'),('Supabase files and bandwidth','1 GB file storage; 5 GB egress','Limit attachments to 5 MB initially; avoid video and large installers'),('Supabase availability','Free projects can pause after a week of inactivity; automatic backups excluded','Use an institute-owned backup process and an availability fallback'),('Firebase Cloud Messaging','Cloud Messaging is listed as no cost','Use FCM only for push; sender compute runs in the Supabase backend'),('Cloudflare Pages','Free static hosting subject to platform limits','Deploy a Vite build to a provider subdomain; a custom domain is optional')],[1.5,2.8,2.4])
p('Realtime connections, message traffic and Edge Function invocation allowances must also be checked against the current Supabase plan before launch. Estimate storage from chat volume and attachments, review usage monthly and restrict uploads before capacity is exhausted. No paid upgrade or billing-enabled service should be enabled automatically.')
h('Email and onboarding')
p('Supabase’s default mail service is intended for limited testing, not routine production delivery. Use institution-provided SMTP or a suitable email provider within its verified free allowance for invites and password reset emails. Confirm that itunit1 and similar addresses are full deliverable mailboxes. If they are shared, preserve individual account ownership and agree a safe reset process. Do not promise free production email without that check.')
h('Cloud deployment decision')
p('The institute should approve external cloud storage of employee contacts, chat and operational records before launch. If cloud hosting is unsuitable, deploy on an institute-managed server and retain the same task rules and SQL data model. Browser push still requires suitable internet connectivity and secure HTTPS access. Free cloud service is a practical pilot choice, but it does not provide a guaranteed production uptime commitment.')
h('Maintenance ownership')
p('The administrator owns account lifecycle, configuration, backups, restore drills and incident handling. The SDD owns assignment policy, report review and business decisions. Keep deployment credentials under institute ownership so the system can be maintained when interns or other staff leave.')

page('7 Delivery plan and acceptance')
p('Indicative duration: 8 weeks for a developer familiar with the selected stack, with weekly feedback from the SDD and representative employees. This is a planning estimate; existing skills and review availability may change it.')
table(['Period','Work','Reviewable outcome'],[
('Week 1','Confirm workflows, directory, permissions and cloud approval','Requirements and screen sketches'),('Week 2','Schema, login, account administration and directory','Secure accounts and team profiles'),('Weeks 3 and 4','Task assignment, progress, completion, history and reports','Working task workflow and sample monthly report'),('Week 5','Direct chat, IT group chat and meeting scheduling','Communication features'),('Week 6','Notification inbox, FCM, reminders and retries','Push pilot on actual institute devices'),('Week 7','Permission testing, report checks, backup and restore drill','Acceptance evidence and defect fixes'),('Week 8','Staff pilot, training, deployment and handover','Production release and maintenance guide')],[1.15,3.0,2.55])
h('Release acceptance criteria')
bullets(['Each member signs in and sees the correct profile, assignments and meetings. SDD and administrator can manage directory records; other employees cannot.','Only the assignee, SDD or administrator override can complete a task, including when using direct API requests. Reassignment and account deactivation take effect immediately.','Completed work and later reopenings appear correctly in monthly reports, including a month-boundary test in Sri Lanka time.','Direct chat is inaccessible to nonparticipants; the IT group works for active members. Administrator moderation actions leave an audit record.','Push works on supported institute devices. Denied permission, expired tokens and failed sends leave a usable notification inbox and retry state.','Meeting updates and cancellations notify invitees. Stored times display correctly in Asia/Colombo.','A backup restores task data and attachment access in a test environment; privileged account and deletion actions are auditable.'])
h('Decisions to confirm before development')
p('Confirm the official institute name and spelling; the supplied name is Sri Lanka Institute of Technology. Confirm Mr. Sumeshan’s designation, all staff spellings, mailbox and LAN extension mappings, the administrator owner, whether employees may view one another’s tasks, and the retention and cloud hosting policies. The completion rule in this proposal allows either the SDD or the assignee to finish a task; confirm if joint approval is preferred.')

page('8 Technical references')
p('Official provider references checked on 8 October 2026. These links support the service plan and technical implementation; requirements and delivery estimates are proposed for this project.')
refs=[('Supabase pricing and free plan allowances','https://supabase.com/pricing'),('Supabase project pausing','https://supabase.com/docs/guides/platform/free-project-pausing'),('Supabase scheduling Edge Functions','https://supabase.com/docs/guides/functions/schedule-functions'),('Supabase SMTP requirements','https://supabase.com/docs/guides/auth/auth-smtp'),('Firebase pricing including Cloud Messaging','https://firebase.google.com/pricing'),('Firebase web messaging setup','https://firebase.google.com/docs/cloud-messaging/web/get-started'),('Cloudflare Pages platform limits','https://developers.cloudflare.com/pages/platform/limits/'),('Cloudflare Pages Functions pricing','https://developers.cloudflare.com/pages/functions/pricing/')]
for title,url in refs: h(title); p(url)
h('Proposed handover package')
p('Deliver the source repository, versioned database migrations and permission policies, deployment instructions, environment variable inventory without secret values, administrator guide, staff quick-start guide, backup and restore procedure, and acceptance test results. Use institute-owned cloud accounts and identify the employee responsible for ongoing maintenance.')
d.core_properties.title='IT Division Work Management System Project Proposal'
d.core_properties.subject='Requirements and implementation plan'
d.core_properties.author=''
d.save('F:/Projects/SLSI/proposal/IT_Division_Project_Proposal.docx')
